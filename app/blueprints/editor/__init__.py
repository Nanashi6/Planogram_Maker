from collections import defaultdict
from typing import Dict, List, Type
from flask import Blueprint, render_template, jsonify, request, send_file
from DataLayer.dao import ProductDAO, ShelfUnitDAO, PlanogramDAO, PlacedProductDAO, RuleDAO, ShelfRuleDAO, CategoryRuleDAO, CategoryDAO
from DataLayer.shemas import Planogram, PlacedProduct, Category, Rule, ShelfRule, CategoryRule
from DataLayer.enums import *
from app.blueprints.editor.exceptions import AutoPlacementError, CommandError
from .models import rules_data, server_message
from .commands_handlers import commands_handler
from .auto_placement import calculate_auto_placement
from .sort_dictionaries import *

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BASE_URL = 'editor'
editor_bp = Blueprint(BASE_URL, __name__, static_folder='static', template_folder='templates', url_prefix=f'/{BASE_URL}')

@editor_bp.route('', methods=['GET'])
@editor_bp.route('/', methods=['GET'])
async def index():
    return render_template(f'{BASE_URL}/index.html')

@editor_bp.route('/get_products', methods=['GET'])
async def get_products():
    products = [product.to_dict() for product in ProductDAO.get_all()]
    return jsonify(products)

@editor_bp.route('/get_shelfUnits', methods=['GET'])
async def get_shelfUnits():
    shelfUnits = [shelfUnit.to_dict() for shelfUnit in ShelfUnitDAO.get_all()]
    return jsonify(shelfUnits)

@editor_bp.route('/get_planograms', methods=['GET'])
async def get_planograms():
    planograms = [planogram.to_dict() for planogram in PlanogramDAO.get_all()]
    return jsonify(planograms)

def enum_to_select_options(enum_class: Type[enum.Enum], custom_labels: Dict[str, str] = None) -> List[Dict[str, str]]:
    """
    Преобразует Enum в список словарей для использования в HTML select.
    Каждый словарь имеет ключи "value" и "label".
    """
    options = []
    if custom_labels is None:
        custom_labels = {}

    for member in enum_class:
        default_label = member.name.replace('_', ' ')
        if default_label.endswith(" Asc"):
            default_label = f"{default_label[:-4]} (По возрастанию / А-Я)"
        elif default_label.endswith(" Desc"):
            default_label = f"{default_label[:-5]} (По убыванию / Я-А)"
        
        label = custom_labels.get(member.value, default_label)
        
        options.append({"value": member.value, "label": label})
    return options

@editor_bp.route('/get_sorting_rules', methods=['GET'])
async def get_sorting_rules():
    brand_sorting_options = enum_to_select_options(BrandSorting, b_sort_keys)
    product_sorting_options = enum_to_select_options(ProductSorting, p_sort_keys)

    sorting_rules = {
        "brand_sort_options": brand_sorting_options,
        "product_sort_options": product_sorting_options,
        "defaults": {
            "brand_sort": BrandSorting.RatingDesc.value,
            "product_sort": ProductSorting.PriceAsc.value
        }
    }
    return jsonify(sorting_rules)

def is_correct_planogram(pps, rules : rules_data):
    errors : List[str] = []

    shelf_rules = rules.shelves_rules
    categoryAllocationRules = defaultdict(list)
    shelf_numbers = {}

    for sr in shelf_rules:
        shelfId = sr.shelf_id
        shelf_numbers[shelfId] = sr.shelf_number
        for cr in sr.category_rules:
            categoryAllocationRules[shelfId].append(cr.category_name)

    for pp in pps:
        product = ProductDAO.get_by_id(pp.get('product_id'))
        if product.category.name not in categoryAllocationRules[pp.get('shelf_id')]:
            errors.append(f'Товар \'{product.name}\' не может быть размещён на полке. Т.к. категория \'{product.category.name}\' отсутствует среди указанных категорий для полки №{shelf_numbers[pp.get('shelf_id')]}')

    is_correct = False if len(errors) > 0 else True
    return is_correct, errors

def save_rule_for_planogram(planogram_id : int, rules : rules_data):
    existing_rule = RuleDAO.get_one(Rule(planogram_id=planogram_id))
    if existing_rule:
        RuleDAO.delete_by_id(existing_rule.id)

    global_rule = Rule(brand_sorting=rules.brand_sort, product_sorting=rules.product_sort, spacing = rules.spacing, planogram_id=planogram_id)
    global_rule_id = RuleDAO.add(global_rule).id

    for sr in rules.shelves_rules:
        sr_id = ShelfRuleDAO.add(ShelfRule(rule_id=global_rule_id, shelf_id=sr.shelf_id)).id
        cat_rules : list[CategoryRule] = []
        for cr in sr.category_rules:
            cat_rules.append(CategoryRule(share = cr.percentage, min_weight=cr.min_weight, max_weight=cr.max_weight, shelf_rule_id=sr_id, category_id=cr.category_id))
        CategoryRuleDAO.add_many(cat_rules)

@editor_bp.route('/save_planogram', methods=['POST'])
async def save_planogram():
    data = request.get_json()
    if not data:
        return jsonify({"error": "No data provided"}), 400

    planogram_data = data.get('planogramData')

    planogram_id = planogram_data.get('id')
    name = planogram_data.get('name')
    shelf_unit_id = planogram_data.get('shelf_unit_id')
    placed_products_payload = planogram_data.get('placed_products', [])

    if not name or shelf_unit_id is None:
        return jsonify({"error": "Missing name or shelf_unit_id"}), 400
    
    rules = rules_data.from_dict(data.get('rules'))

    is_correct, error_messages = is_correct_planogram(placed_products_payload, rules)
    if is_correct:
        new_planogram = Planogram(name = name, shelf_unit_id = shelf_unit_id)
        try:
            if planogram_id: # Обновление существующей планограммы
                planogram = PlanogramDAO.get_by_id(planogram_id)
                if not planogram:
                    return jsonify({"error": f"Planogram with id {planogram_id} not found"}), 404
                PlanogramDAO.update_by_id(planogram_id, new_planogram)
                PlacedProductDAO.delete_many(PlacedProduct(planogram_id=planogram_id))
            else: # Создание новой планограммы
                planogram = PlanogramDAO.add(new_planogram)
                if not planogram or not planogram.id:
                    return jsonify({"error": "Failed to create new planogram"}), 500
                planogram_id = planogram.id

            new_pps = [PlacedProduct(shelf_id=pp.get('shelf_id'), product_id=pp.get('product_id'), position=pp.get('position'), planogram_id=planogram_id) for pp in placed_products_payload]
            pps = PlacedProductDAO.add_many(new_pps)

            final_planogram_obj = PlanogramDAO.get_by_id(planogram_id)
            if not final_planogram_obj:
                return jsonify({"error": "Failed to retrieve saved planogram details"}), 500

            # Беру правила и сохраняю их с привязкой к планограмме + правила для полок + правила для категорий
            save_rule_for_planogram(planogram_id, rules)

            return jsonify(final_planogram_obj.to_dict()), 200

        except Exception as e:
            print(f"Error saving planogram: {e}")
            return jsonify({"error": str(e)}), 500
    else:
        return jsonify({'errors': error_messages}), 200

@editor_bp.route('/chat_message', methods=['POST'])
async def message_handle():
    try: 
        data = request.get_json()
        if not data:
            return jsonify({"message": "Ошибка: Тело запроса не содержит JSON."}), 400
        
        user_message = data.get('message')
        planogram_data = data.get('planogram')
        rules_data_dict = data.get('rules') 

        if not user_message:
            return jsonify({"message": "Ошибка: Ключ 'message' отсутствует или пуст."}), 400
        
        success_response_message = commands_handler(user_message, planogram_data, rules_data_dict)
        return jsonify(success_response_message.to_json()), 200

    except CommandError as e:
        print(f"Command Error: {e}")
        error_response = server_message(
            message=str(e),
            data=data.get('planogram'),
            rules=rules_data.from_dict(data.get('rules')) 
        )
        return jsonify(error_response.to_json()), 200

    except Exception as e:
        import traceback
        print("Critical error in /chat_message endpoint:")
        print(e)
        traceback.print_exc()
        critical_error_response = server_message(
            message="Внутренняя ошибка сервера. Не удалось обработать команду.",
            data=request.get_json().get('planogram'),
            rules=rules_data.from_dict(request.get_json().get('rules'))
        )
        return jsonify(critical_error_response.to_json()), 200

@editor_bp.route('/process_rules_file', methods=['POST'])
async def process_rules_file():
    if 'rules_file' not in request.files:
        return jsonify({"error": "Файл не найден в запросе."}), 400

    file = request.files['rules_file']
    if file.filename == '':
        return jsonify({"error": "Файл не выбран."}), 400

    shelf_unit_id = request.form.get('shelf_unit_id', type=int)
    if not shelf_unit_id:
        return jsonify({"error": "Не указан ID стеллажа для применения правил."}), 400

    shelf_unit = ShelfUnitDAO.get_by_id(shelf_unit_id)
    if not shelf_unit:
        return jsonify({"error": f"Стеллаж с ID {shelf_unit_id} не найден."}), 404

    try:
        file_content = file.read().decode('utf-8')
        commands = file_content.splitlines()

        current_rules = rules_data()
        current_planogram_dict = {'shelf_unit_id': shelf_unit.id, 'shelf_unit': shelf_unit.to_dict()}
        print(current_planogram_dict)
        for i, command_str in enumerate(commands):
            command_str = command_str.strip()
            if not command_str or command_str.startswith('#'):
                continue 
            print(command_str)
            response_message = commands_handler(command_str, current_planogram_dict, current_rules.to_json())
            current_planogram_dict = response_message.data.to_json()
            print(current_planogram_dict)
            current_rules = response_message.rules
        
        
        return jsonify(current_rules.to_json()), 200

    except CommandError as e:
        line_num = i + 1 if 'i' in locals() else 'N/A'
        return jsonify({"error": f"Ошибка в файле на строке {line_num}: {e}"}), 400
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": f"Произошла внутренняя ошибка сервера при обработке файла: {e}"}), 500
    
@editor_bp.route('/calculate_auto_placement', methods=['POST'])
def calculate_auto_placement_route():
    try: 
        data = request.get_json()
        if not data:
            return jsonify({"message": "Ошибка: Тело запроса не содержит JSON."}), 400
        
        rules = data.get('rules')
        shelf_unit_id = data.get('shelf_unit_id')
        planogram_id = data.get('planogram_id')
        planogram_name = data.get('planogram_name')

        if not rules or not shelf_unit_id:
            return jsonify({"message": "Ошибка: Запрос не содержит правила или ID стеллажа."}), 400

        server_message = calculate_auto_placement(shelf_unit_id, planogram_id, planogram_name, rules_data.from_dict(rules))
        return jsonify(server_message.to_json()), 200

    except AutoPlacementError as e:
        print(f"Ошибка авторасстановки: {e}")
        return jsonify({"message": str(e)}), 400
    except Exception as e:
        print(f"Критическая ошибка в /calculate_auto_placement: {e}")
        # traceback.print_exc()
        return jsonify({
            "message": "Внутренняя ошибка сервера. Пожалуйста, обратитесь к администратору."
        }), 500

@editor_bp.route('/download_xlsx', methods=['POST'])
def download_xlsx():
    try:
        planogram_data = request.get_json()
        if not planogram_data:
            return jsonify({"error": "No data provided"}), 400

        workbook = Workbook()
        sheet = workbook.active
        
        planogram_name = planogram_data.get('planogramName', 'Планограмма')
        sheet.title = planogram_name[:31]

        shelves_content = {}
        sorted_shelves = sorted(planogram_data.get('shelves', []), key=lambda s: s.get('shelfNumber', 0))

        for shelf in sorted_shelves:
            shelf_title = f"Полка {shelf.get('shelfNumber', 'N/A')}"
            products_on_shelf = []
            
            all_products = []
            for category in shelf.get('categories', []):
                all_products.extend(category.get('products', []))
            
            sorted_products = sorted(all_products, key=lambda p: p.get('positionOnShelf', 0))
            
            products_on_shelf = [p.get('name', 'Без названия') for p in sorted_products]
            shelves_content[shelf_title] = products_on_shelf
        
        headers = list(shelves_content.keys())
        for col_num, header_text in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_num, value=header_text)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

        max_rows = 0
        if shelves_content:
            max_rows = max(len(products) for products in shelves_content.values())

        for col_num, shelf_title in enumerate(headers, 1):
            products = shelves_content.get(shelf_title, [])
            for row_num, product_name in enumerate(products, 2):
                sheet.cell(row=row_num, column=col_num, value=product_name)

        for col_num, _ in enumerate(headers, 1):
            column_letter = get_column_letter(col_num)
            max_length = 0
            for i in range(1, max_rows + 2): 
                cell_value = sheet.cell(row=i, column=col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            adjusted_width = (max_length + 2)
            sheet.column_dimensions[column_letter].width = adjusted_width

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        filename = f"{planogram_name.replace(' ', '_')}.xlsx"
        return send_file(
            buffer,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": f"Произошла внутренняя ошибка сервера: {str(e)}"}), 500

# # TODO Учитывать доли брендов на полках